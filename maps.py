"""
Поиск местных бизнесов по OpenStreetMap -> Excel с приоритетами -> твой бот.

Для каждой точки: ниша, адрес, контакты, есть ли сайт и что с ним не так,
что ей предложить, ссылки на Яндекс.Карты и 2ГИС для проверки.

  python maps.py            # обычный прогон (шлёт Excel в бота)
  python maps.py --dry-run  # только собрать leads.xlsx локально
"""
import datetime as dt
import json
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import quote, urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common import CONFIG, HEADERS, ROOT, load_json, now_iso, save_json, send_document

CFG = CONFIG["maps"]
OUT_FILE = ROOT / "leads.xlsx"
TODAY = dt.date.today().isoformat()
OVERPASS = ["https://overpass-api.de/api/interpreter",
            "https://overpass.kumi.systems/api/interpreter",
            "https://maps.mail.ru/osm/tools/overpass/api/interpreter"]

SOCIAL = ("vk.com", "vk.ru", "instagram.com", "t.me", "telegram.me",
          "ok.ru", "wa.me", "facebook.com", "youtube.com", "dzen.ru", "max.ru")
# шаблонные странички вместо сайта: мультиссылки, QR-меню, бесплатные адреса конструкторов
TEMPLATE = ("taplink.", "linktr.ee", "mssg.me", "hipolink", "choiceqr", "qr-menu", "qrmenu", "menu.qr",
            "tilda.ws", "wixsite.com", "ucoz.", "narod.ru", "business.site", "site123.me", "jimdosite",
            "nethouse.ru/", "mya5.ru", "tb.ru", "yandex.ru/sprav", "sprav.", "ukit.me", "webflow.io")
BOOKING = ("yclients", "dikidi", "sonline", "booksy", "arnica")
MAPS_HOSTS = ("yandex.", "2gis.", "google.", "goo.gl")
BUILDERS = {"tilda": "Tilda", "wix.com": "Wix", "nethouse": "Nethouse",
            "ukit": "uKit", "mottor": "Mottor", "1c-bitrix": "Битрикс",
            "wordpress": "WordPress", "flexbe": "Flexbe", "lpgenerator": "LPgenerator"}

MENU_WORDS = r"меню|menu|барная карта|винная карта|бар-карта"
MENU_FILE = r"href=[\"']([^\"']*(?:menu|menyu|меню)[^\"']*\.(?:pdf|jpe?g|png|webp))"

PRIO_LABEL = {3: "🔥 Горячий", 2: "Тёплый", 1: "Слабый", 0: "Сайт в порядке"}


# ---------- теги и ниши ----------

def parse_tag(s):
    k, _, v = s.partition("=")
    return k.strip(), [x.strip() for x in v.split("|")] if v else []


def tag_ok(tags, spec):
    k, vals = parse_tag(spec)
    return k in tags and (not vals or tags[k] in vals)


def classify(tags):
    name = (tags.get("name") or "").lower()
    for n in CFG["niches"]:
        if any(tag_ok(tags, s) for s in n.get("tags", [])):
            return n
        bn = n.get("by_name")
        if bn and any(tag_ok(tags, s) for s in bn["tags"]) and re.search(bn["words"], name):
            return n
    return None


def all_specs():
    specs = []
    for n in CFG["niches"]:
        specs += n.get("tags", [])
        if n.get("by_name"):
            specs += n["by_name"]["tags"]
    return sorted(set(specs))


def build_query(region):
    q = [f'[out:json][timeout:600];',
         f'area["name"="{region["name"]}"]["boundary"="administrative"]'
         f'["admin_level"="{region.get("admin_level", 4)}"]->.city;']
    if region.get("districts"):
        names = "|".join(re.escape(d) for d in region["districts"])
        q.append(f'rel(area.city)["boundary"="administrative"]["name"~"^({names})$"];')
        q.append("map_to_area->.a;")
    else:
        q.append(".city->.a;")
    q.append("(")
    for spec in all_specs():
        k, vals = parse_tag(spec)
        if not vals:
            q.append(f'  nwr["{k}"](area.a);')
        elif len(vals) == 1:
            q.append(f'  nwr["{k}"="{vals[0]}"](area.a);')
        else:
            q.append(f'  nwr["{k}"~"^({"|".join(vals)})$"](area.a);')
    q.append(");")
    q.append("out center tags;")
    return "\n".join(q)


def fetch_osm(region):
    query = build_query(region)
    last = None
    for url in OVERPASS:
        for attempt in range(2):
            try:
                r = requests.post(url, data={"data": query}, timeout=700,
                                  headers={"User-Agent": "lead-parser/1.0 (small business research)"})
                if r.status_code in (429, 504):
                    time.sleep(30)
                    continue
                r.raise_for_status()
                return r.json()["elements"]
            except Exception as ex:
                last = ex
                print(f"  [{urlparse(url).netloc}] ошибка: {ex}")
                time.sleep(10)
    raise RuntimeError(f"Overpass недоступен: {last}")


# ---------- сайты ----------

def host_of(url):
    return urlparse(url if "://" in url else "http://" + url).netloc.lower()


def kind_of(url):
    h = host_of(url) + urlparse(url if "://" in url else "http://" + url).path.lower()
    if any(s in h for s in SOCIAL):
        return "social"
    if any(s in h for s in BOOKING):
        return "booking"
    if any(s in h for s in TEMPLATE):
        return "template"
    if any(s in h for s in MAPS_HOSTS):
        return "maps"
    return "site"


def check_site(url):
    full = url if "://" in url else "http://" + url
    try:
        t0 = time.time()
        r = requests.get(full, headers=HEADERS, timeout=(7, 15), allow_redirects=True)
        elapsed = time.time() - t0
    except Exception as ex:
        return {"ok": False, "err": type(ex).__name__}
    if r.status_code >= 400:
        return {"ok": False, "err": f"HTTP {r.status_code}"}
    body = r.content[:500_000].decode(r.encoding or "utf-8", errors="ignore") \
        if "html" in r.headers.get("content-type", "html") else ""
    low = body.lower()
    years = [int(y) for y in re.findall(
        r"(?:©|&copy;|copyright)\s*(?:\d{4}\s*(?:-|–|—)\s*)?(20\d\d)", low)]
    return {
        "ok": True,
        "https": r.url.startswith("https://"),
        "viewport": bool(re.search(r"<meta[^>]+name=[\"']?viewport", low)),
        "slow": elapsed > 5,
        "year": max(years) if years else None,
        "menu": menu_status(low),
        "builder": next((v for k, v in BUILDERS.items() if k in low), ""),
        "parked": len(low) < 1500 or any(w in low for w in (
            "домен продается", "domain is for sale", "this domain", "сайт находится в разработке",
            "under construction", "hosting is suspended", "срок регистрации")),
    }


def menu_status(low):
    if re.search(MENU_FILE, low):
        return "Меню файлом (PDF/картинка)"
    # слово «меню» в ссылках/заголовках, а не только в навигации сайта «Меню ☰»
    if re.search(r"<a[^>]+href=[\"'][^\"']*(?:menu|menyu|меню|food|eda|kitchen|kuhnya)", low) \
            or len(re.findall(MENU_WORDS, low)) >= 3:
        return "Есть на сайте"
    return "Нет на сайте"


def menu_note(lead, site):
    """Что с меню у заведения — для колонки «Меню»."""
    if not lead["menu"]:
        return ""
    if lead["menu_link"]:
        return "Ссылка на меню в OSM — посмотри, как оформлено"
    if lead["site_kind"] != "site" or not lead["site"]:
        return "Сайта нет — меню только на картах/в соцсетях?"
    if not site or not site["ok"] or site.get("parked"):
        return "Проверь вручную"
    return site["menu"]


def flags_of(lead, site):
    """Короткие коды проблем — для фильтров и бейджей в приложении."""
    f = []
    has_contact = any(lead[k] for k in ("phone", "vk", "tg", "insta", "email"))
    if lead["site_kind"] == "social":
        f.append("social")
    elif lead["site_kind"] == "template":
        f.append("template")
    elif lead["site_kind"] == "booking":
        f.append("booking")
    elif not lead["site"]:
        f.append("no_site")
    elif site is not None:
        if not site["ok"]:
            f.append("down")
        elif site["parked"]:
            f.append("parked")
        else:
            if not site["viewport"]:
                f.append("no_mobile")
            if site["year"] and site["year"] <= dt.date.today().year - CFG.get("old_site_years", 3):
                f.append("old")
            if site["slow"]:
                f.append("slow")
            if not site["https"]:
                f.append("no_https")
            if lead["menu"] and site.get("menu") == "Нет на сайте":
                f.append("no_menu")
            if lead["menu"] and site.get("menu", "").startswith("Меню файлом"):
                f.append("menu_file")
    if not has_contact:
        f.append("no_contacts")
    return f


def assess(lead, site):
    """Возвращает (приоритет 0..3, список проблем)."""
    problems, prio = [], 0
    has_contact = any(lead[k] for k in ("phone", "vk", "tg", "insta", "email"))
    if lead["site_kind"] == "social":
        return 3, ["Вместо сайта — соцсеть/таплинк"]
    if lead["site_kind"] == "template":
        return 3, ["Шаблонная страничка вместо сайта (мультиссылка, QR-меню, бесплатный конструктор)"]
    if lead["site_kind"] == "booking":
        return 2, ["Только виджет записи, своего сайта нет"]
    if not lead["site"]:
        if has_contact:
            return 3, ["Нет сайта"]
        return 1, ["Нет сайта и контактов в OSM — найди их на картах"]
    if site is None:
        return 0, ["Сайт не проверялся"]
    if not site["ok"]:
        return 2, [f"Сайт не открылся ({site['err']}) — проверь вручную"]
    if site["parked"]:
        return 3, ["Сайт-заглушка / пустой / домен истёк"]
    if not site["viewport"]:
        prio, problems = 2, problems + ["Нет мобильной версии"]
    if site["year"] and site["year"] <= dt.date.today().year - CFG.get("old_site_years", 3):
        prio, problems = max(prio, 1), problems + [f"Давно не обновлялся (© {site['year']})"]
    if site["slow"]:
        prio, problems = max(prio, 1), problems + ["Грузится дольше 5 сек"]
    if not site["https"]:
        prio, problems = max(prio, 1), problems + ["Нет https"]
    if lead["menu"] and site.get("menu") == "Нет на сайте":
        prio, problems = max(prio, 2), problems + ["На сайте нет меню"]
    elif lead["menu"] and site.get("menu", "").startswith("Меню файлом"):
        prio, problems = max(prio, 1), problems + ["Меню скачивается файлом — можно сделать красивее"]
    if site["builder"]:
        problems.append(f"Конструктор: {site['builder']}")
    return prio, problems or ["Видимых проблем нет"]


# ---------- сборка ----------

def first(tags, *keys):
    for k in keys:
        if tags.get(k):
            return tags[k].split(";")[0].strip()
    return ""


def to_lead(el, niche, region):
    t = el.get("tags", {})
    site = first(t, "website", "contact:website", "url")
    vk = first(t, "contact:vk")
    insta = first(t, "contact:instagram")
    tg = first(t, "contact:telegram")
    kind = kind_of(site) if site else ""
    if kind in ("social", "booking", "maps"):
        h = host_of(site)
        if "vk." in h and not vk:
            vk = site
        elif "instagram" in h and not insta:
            insta = site
        elif ("t.me" in h or "telegram" in h) and not tg:
            tg = site
    if kind == "maps":
        site, kind = "", ""
    addr = ", ".join(x for x in (t.get("addr:street", ""), t.get("addr:housenumber", "")) if x)
    pos = el.get("center") or el
    return {
        "id": f"{el['type']}/{el['id']}",
        "region": region["name"], "twogis": region.get("twogis", "spb"),
        "city": t.get("addr:city") or t.get("addr:place") or region.get("city", ""),
        "lat": pos.get("lat"), "lon": pos.get("lon"),
        "niche": niche["name"], "offer": niche.get("offer", ""),
        "name": t.get("name", ""), "addr": addr,
        "phone": first(t, "phone", "contact:phone", "contact:mobile"),
        "email": first(t, "email", "contact:email"),
        "site": site,
        "site_kind": kind, "vk": vk, "tg": tg, "insta": insta,
        "hours": t.get("opening_hours", ""),
        "menu": bool(niche.get("menu")),
        "menu_link": first(t, "website:menu", "menu:url", "contact:menu"),
    }


def blocked_ids():
    """Лиды со статусом «Блок» в приложении: робот их больше не берёт."""
    crm = load_json("crm.json", {})
    return {i for i, v in (crm.get("leads") or {}).items() if isinstance(v, dict) and v.get("s") == "block"}


def collect():
    leads, done = {}, []
    blocked = blocked_ids()
    skipped = 0
    for region in CFG["regions"]:
        try:
            elements = fetch_osm(region)
        except Exception as ex:
            print(f"[{region['name']}] пропущен: {ex}")
            continue
        before, chains = len(leads), 0
        for el in elements:
            tags = el.get("tags", {})
            if not tags.get("name"):
                continue
            if CFG.get("skip_chains", True) and (tags.get("brand") or tags.get("brand:wikidata")):
                chains += 1
                continue
            if f"{el['type']}/{el['id']}" in blocked:
                skipped += 1
                continue
            niche = classify(tags)
            if niche:
                lead = to_lead(el, niche, region)
                leads.setdefault(lead["id"], lead)
        done.append(region["name"])
        print(f"[{region['name']}] объектов: {len(elements)}, подходящих: "
              f"{len(leads) - before}, сетей пропущено: {chains}")
        time.sleep(5)
    if skipped:
        print(f"В блоке: {skipped}, пропущены")
    return list(leads.values()), done


def check_all(leads):
    targets = [l for l in leads if l["site"] and l["site_kind"] == "site"]
    results = {}
    if CFG.get("check_sites", True) and targets:
        print(f"[сайты] проверяю {len(targets)}…")
        with ThreadPoolExecutor(max_workers=32) as pool:
            for lead, res in zip(targets, pool.map(lambda l: check_site(l["site"]), targets)):
                results[lead["id"]] = res
    for l in leads:
        l["prio"], probs = assess(l, results.get(l["id"]))
        l["problems"] = "; ".join(probs)
        l["menu_note"] = menu_note(l, results.get(l["id"]))
        l["flags"] = flags_of(l, results.get(l["id"]))
        res = results.get(l["id"]) or {}
        l["builder"] = res.get("builder", "")


# ---------- Excel ----------

COLS = [("Приоритет", 14), ("Новый", 8), ("Регион", 20), ("Город", 18), ("Ниша", 20),
        ("Название", 30), ("Адрес", 28),
        ("Проблема", 38), ("Меню", 26), ("Что предложить", 42), ("Телефон", 18), ("Сайт", 28),
        ("VK", 22), ("Telegram", 20), ("Instagram", 20), ("Email", 24),
        ("Часы работы", 22), ("Яндекс.Карты", 14), ("2ГИС", 10),
        ("Впервые найден", 14), ("Статус", 16), ("Заметки", 30)]


def url_of(v):
    if not v:
        return ""
    if v.startswith("@"):
        return ""
    return v if "://" in v else "https://" + v


def map_links(l):
    name = quote(l["name"])
    if l["lat"] and l["lon"]:
        ya = f"https://yandex.ru/maps/?ll={l['lon']},{l['lat']}&z=18&text={name}"
        gis = f"https://2gis.ru/{l['twogis']}/search/{name}?m={l['lon']}%2C{l['lat']}%2F18"
    else:
        where = ", ".join(x for x in (l["name"], l["city"] or l["region"], l["addr"]) if x)
        ya = "https://yandex.ru/maps/?text=" + quote(where)
        gis = f"https://2gis.ru/{l['twogis']}/search/{name}"
    return ya, gis


def col(name):
    return [c for c, _ in COLS].index(name) + 1


def write_xlsx(leads, fresh):
    wb = Workbook()
    ws = wb.active
    ws.title = "Лиды"
    ws.append([c for c, _ in COLS])
    head_fill = PatternFill("solid", fgColor="1F2937")
    for i, (_, w) in enumerate(COLS, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w

    fills = {3: "FDE2E1", 2: "FEF3C7", 1: "F3F4F6"}
    leads.sort(key=lambda l: (-l["prio"], not fresh(l), l["region"], l["niche"], l["name"]))
    for l in leads:
        ya, gis = map_links(l)
        row = [PRIO_LABEL[l["prio"]], "да" if fresh(l) else "", l["region"], l["city"],
               l["niche"], l["name"], l["addr"], l["problems"], l["menu_note"], l["offer"], l["phone"],
               l["site"], l["vk"], l["tg"], l["insta"], l["email"], l["hours"],
               "открыть", "открыть", l["first_seen"], "", ""]
        ws.append(row)
        r = ws.max_row
        links = {col("Меню"): url_of(l["menu_link"]), col("Сайт"): url_of(l["site"]), col("VK"): url_of(l["vk"]),
                 col("Telegram"): url_of(l["tg"]), col("Instagram"): url_of(l["insta"]),
                 col("Яндекс.Карты"): ya, col("2ГИС"): gis}
        for ci, u in links.items():
            if u:
                c = ws.cell(row=r, column=ci)
                c.hyperlink = u
                c.font = Font(color="2563EB", underline="single")
        if l["prio"] in fills:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fills[l["prio"]])

    dv = DataValidation(type="list", allow_blank=True,
                        formula1='"Не связалась,Написала,Ответили,Созвон,Отказ,Клиент"')
    ws.add_data_validation(dv)
    st = get_column_letter(col("Статус"))
    dv.add(f"{st}2:{st}{max(ws.max_row, 2)}")
    ws.freeze_panes = ws.cell(row=2, column=col("Название") + 1)
    ws.auto_filter.ref = ws.dimensions

    # отдельный лист «Меню»: только еда, отсортировано по тому, как сильно им нужно меню
    food = [l for l in leads if l["menu"]]
    if food:
        order = {"Сайта нет — меню только на картах/в соцсетях?": 0, "Нет на сайте": 1,
                 "Меню файлом (PDF/картинка)": 2, "Проверь вручную": 3}
        m = wb.create_sheet("Меню")
        mcols = [("Меню", 34), ("Регион", 20), ("Город", 18), ("Ниша", 20), ("Название", 30),
                 ("Адрес", 28), ("Телефон", 18), ("Сайт", 28), ("VK", 22), ("Telegram", 20),
                 ("Яндекс.Карты", 14), ("2ГИС", 10), ("Статус", 16), ("Заметки", 30)]
        m.append([c for c, _ in mcols])
        for i, (_, w) in enumerate(mcols, 1):
            cell = m.cell(row=1, column=i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            m.column_dimensions[get_column_letter(i)].width = w
        for l in sorted(food, key=lambda l: (order.get(l["menu_note"], 4), l["region"], l["name"])):
            ya, gis = map_links(l)
            m.append([l["menu_note"], l["region"], l["city"], l["niche"], l["name"], l["addr"],
                      l["phone"], l["site"], l["vk"], l["tg"], "открыть", "открыть", "", ""])
            r = m.max_row
            for ci, u in {8: url_of(l["site"]), 9: url_of(l["vk"]), 10: url_of(l["tg"]),
                          11: ya, 12: gis}.items():
                if u:
                    m.cell(row=r, column=ci).hyperlink = u
                    m.cell(row=r, column=ci).font = Font(color="2563EB", underline="single")
        mdv = DataValidation(type="list", allow_blank=True,
                             formula1='"Не связалась,Написала,Ответили,Созвон,Отказ,Клиент"')
        m.add_data_validation(mdv)
        mdv.add(f"M2:M{max(m.max_row, 2)}")
        m.freeze_panes = "F2"
        m.auto_filter.ref = m.dimensions

    # сводка
    s = wb.create_sheet("Сводка")
    s.append(["Регион", "Ниша", "🔥 Горячих", "Тёплых", "Слабых", "Сайт в порядке", "Всего", "Новых"])
    for c in s[1]:
        c.font = Font(bold=True)
    for region in [r["name"] for r in CFG["regions"]]:
        for niche in [n["name"] for n in CFG["niches"]]:
            group = [l for l in leads if l["niche"] == niche and l["region"] == region]
            if group:
                s.append([region, niche] + [sum(l["prio"] == p for l in group) for p in (3, 2, 1, 0)] +
                         [len(group), sum(fresh(l) for l in group)])
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 26
    for c in "CDEFGH":
        s.column_dimensions[c].width = 13
    wb.save(OUT_FILE)


def main():
    dry = "--dry-run" in sys.argv
    state = load_json("maps_state.json", {})
    if "ids" not in state:
        state = {"ids": {}, "regions": []}
    known = set(state["regions"])  # регионы, которые уже собирались раньше

    leads, done = collect()
    if not leads:
        raise SystemExit("Ничего не собрано — карты недоступны, попробуй позже.")
    check_all(leads)
    for l in leads:
        l["first_seen"] = state["ids"].setdefault(l["id"], TODAY)
    fresh = lambda l: l["region"] in known and l["first_seen"] == TODAY
    write_xlsx(leads, fresh)
    state["regions"] = sorted(known | set(done))
    save_json("maps_state.json", state, compact=True)

    # для приложения; регионы, которые сейчас не собрались, берём из прошлого раза
    prev = load_json("leads.json", {"items": []})
    blocked = blocked_ids()
    keep = [l for l in prev["items"] if l["region"] not in done and l["id"] not in blocked]
    fields = ("id", "region", "city", "niche", "name", "addr", "phone", "email", "site",
              "site_kind", "vk", "tg", "insta", "hours", "lat", "lon", "prio", "problems",
              "flags", "menu", "menu_note", "menu_link", "offer", "builder", "first_seen")
    items = [{k: l.get(k) for k in fields} for l in leads]
    for it in items:
        it["new"] = fresh(it)
    save_json("leads.json", {"updated": now_iso(), "regions": sorted(set(done) | {l["region"] for l in keep}),
                             "items": items + keep}, compact=True)

    hot = sum(l["prio"] == 3 for l in leads)
    warm = sum(l["prio"] == 2 for l in leads)
    new = sum(fresh(l) for l in leads)
    caption = (f"🗺 <b>Лиды с карт</b>\n{', '.join(done)}\n"
               f"Всего: {len(leads)} · 🔥 горячих: {hot} · тёплых: {warm}"
               + (f"\nНовых с прошлого раза: {new}" if known else ""))
    print(caption.replace("<b>", "").replace("</b>", ""))
    if not dry:
        send_document(OUT_FILE, caption)


if __name__ == "__main__":
    main()
